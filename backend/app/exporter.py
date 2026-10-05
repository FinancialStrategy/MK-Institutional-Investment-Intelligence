from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import SimulationInput
from .tax_registry import public_registry


NAVY = "0A1630"
ORANGE = "F39A2E"
WHITE = "FFFFFF"
LIGHT = "E8EDF5"
DARK_TEXT = "17233A"
INPUT_BLUE = "0066FF"
GREEN = "1E8F62"
GRAY = "6F7D91"


LABELS = {
    "en": {
        "title": "INTEREST LOAN COST SIMULATOR INTELLIGENCE",
        "summary": "Executive Summary",
        "inputs": "Input Snapshot",
        "tax": "Tax Registry",
        "borrower": "Borrower Cost",
        "investment": "Investment Return",
        "carry": "Net Carry & Break-Even",
        "lender": "Lender Economics",
        "curve": "Loan Yield Curve",
        "ladder": "Maturity Ladder",
        "scenarios": "Scenarios",
        "sensitivity": "Sensitivity",
        "methodology": "Methodology",
        "sources": "Sources",
        "real": "Inflation & Real Cost",
    },
    "tr": {
        "title": "INTEREST LOAN COST SIMULATOR INTELLIGENCE",
        "summary": "Y\u00f6netici \u00d6zeti",
        "inputs": "Girdi Anl\u0131k G\u00f6r\u00fcn\u00fcm\u00fc",
        "tax": "Vergi Sicili",
        "borrower": "Kredi Kullanan Maliyeti",
        "investment": "Yat\u0131r\u0131m Getirisi",
        "carry": "Net Carry ve Ba\u015faba\u015f",
        "lender": "Kredi Veren Ekonomisi",
        "curve": "Kredi Faizi Verim E\u011frisi",
        "ladder": "Vade Merdiveni",
        "scenarios": "Senaryolar",
        "sensitivity": "Duyarl\u0131l\u0131k",
        "methodology": "Metodoloji",
        "sources": "Kaynaklar",
        "real": "Enflasyon ve Reel Maliyet",
    },
}


def _style_sheet(ws, title: str):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:H2")
    c = ws["A1"]
    c.value = title
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(name="Arial", size=16, bold=False, color=WHITE)
    c.alignment = Alignment(vertical="center", horizontal="left")
    ws.merge_cells("A3:H3")
    ws["A3"] = "MK FinTECH LabGEN @2026 Istanbul Atelier"
    ws["A3"].font = Font(name="Arial", size=9, bold=False, color=ORANGE)
    ws["A3"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A3"].alignment = Alignment(horizontal="left")
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 20


def _write_pairs(ws, start_row: int, pairs: list[tuple[str, object]], percent_keys: set[str] | None = None):
    percent_keys = percent_keys or set()
    thin = Side(style="thin", color="D6DCE5")
    for i, (label, value) in enumerate(pairs, start_row):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value)
        ws.cell(i, 1).font = Font(name="Arial", size=10, bold=False, color=DARK_TEXT)
        ws.cell(i, 2).font = Font(name="Arial", size=10, bold=False, color=DARK_TEXT)
        ws.cell(i, 1).border = Border(bottom=thin)
        ws.cell(i, 2).border = Border(bottom=thin)
        if label in percent_keys and isinstance(value, (float, int)):
            ws.cell(i, 2).number_format = "0.00%"
        elif isinstance(value, (float, int)):
            ws.cell(i, 2).number_format = '#,##0.00;[Red](#,##0.00);-'


def build_results_workbook(inp: SimulationInput, result: dict, language: str) -> bytes:
    t = LABELS[language]
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(t["summary"])
    _style_sheet(ws, t["title"])
    k = result["kpis"]
    summary_pairs = [
        ("TLREF Effective Annual", k["tlref_effective_annual"]),
        ("All-in Borrowing Effective Annual", k["all_in_borrowing_effective_annual"]),
        ("After-tax Borrowing Effective Annual", k["after_tax_borrowing_effective_annual"]),
        ("After-tax Investment Effective Annual", k["after_tax_investment_effective_annual"]),
        ("Net Carry Amount", k["net_carry_amount"]),
        ("Net Carry Effective Annual", k["net_carry_effective_annual"]),
        ("Break-even Loan Rate", k["break_even_loan_rate"]),
        ("Break-even Spread", k["break_even_spread"]),
        ("Safety Margin", k["safety_margin"]),
        ("10Y Bond Relative Value", k["ten_year_bond_relative_value"]),
        ("P(Net Carry > 0)", result["scenarios"]["positive_carry_probability"]),
        ("Expected Net Carry Amount", result["scenarios"]["expected_net_carry_amount"]),
    ]
    pct = {x[0] for x in summary_pairs if "Amount" not in x[0]}
    _write_pairs(ws, 5, summary_pairs, pct)

    ws = wb.create_sheet(t["inputs"])
    _style_sheet(ws, t["inputs"])
    input_pairs = [
        ("Principal", inp.principal), ("Valuation Date", inp.valuation_date.isoformat()), ("Days", inp.days), ("Curve Mode", inp.curve_mode), ("TLREF", inp.tlref),
        ("Direct Loan Rate", inp.direct_loan_rate), ("Credit Spread", inp.credit_spread),
        ("12M Expected CPI", inp.inflation_12m_expectation), ("Scenario Inflation", inp.inflation), ("10Y TL Bond Yield", inp.ten_year_bond_yield),
        ("BSMV", inp.bsmv_rate), ("KKDF", inp.kkdf_rate), ("Withholding", inp.withholding_rate),
        ("Corporate Tax", inp.corporate_tax_rate), ("Lender Tax", inp.lender_tax_rate),
        ("Deductible Ratio", inp.deductible_ratio), ("Upfront Fee", inp.upfront_fee_rate),
        ("Commission", inp.commission_rate), ("Lender Funding", inp.lender_funding_rate),
        ("Lender ECL", inp.lender_ecl_rate), ("Lender Opex", inp.lender_opex_rate),
        ("Lender Capital/Liquidity", inp.lender_capital_rate),
    ]
    _write_pairs(ws, 5, input_pairs, {x[0] for x in input_pairs if x[0] not in {"Principal", "Days"}})

    ws = wb.create_sheet(t["tax"])
    _style_sheet(ws, t["tax"])
    tax_headers = ["Key", "Rate", "Effective Date", "Source", "Description"]
    for c, h in enumerate(tax_headers, 1):
        ws.cell(5, c, h)
        ws.cell(5, c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(5, c).font = Font(name="Arial", size=9, bold=False, color=WHITE)
    for r, item in enumerate(public_registry(), 6):
        desc = item["label_tr"] if language == "tr" else item["label_en"]
        vals = [item["key"], item["rate"], item["effective_date"], item["source"], desc]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = Font(name="Arial", size=9, bold=False, color=DARK_TEXT)
        ws.cell(r, 2).number_format = "0.00%"

    for key, sheet_name, section in [
        ("borrower", t["borrower"], result["borrower"]),
        ("investment", t["investment"], result["investment"]),
        ("lender", t["lender"], result["lender"]),
    ]:
        ws = wb.create_sheet(sheet_name)
        _style_sheet(ws, sheet_name)
        pairs = [(k.replace("_", " ").title(), v) for k, v in section.items()]
        pct_labels = {label for label, value in pairs if isinstance(value, float) and abs(value) <= 5 and any(x in label.lower() for x in ["rate", "spread", "effective", "margin"])}
        _write_pairs(ws, 5, pairs, pct_labels)

    ws = wb.create_sheet(t["carry"])
    _style_sheet(ws, t["carry"])
    carry_pairs = [
        ("Net Carry Amount", k["net_carry_amount"]),
        ("Net Carry Period Rate", k["net_carry_period_rate"]),
        ("Net Carry Effective Annual", k["net_carry_effective_annual"]),
        ("Break-even Loan Rate", k["break_even_loan_rate"]),
        ("Break-even Spread", k["break_even_spread"]),
        ("Safety Margin", k["safety_margin"]),
        ("10Y Bond Relative Value", k["ten_year_bond_relative_value"]),
        ("Expected Net Carry Amount", result["scenarios"]["expected_net_carry_amount"]),
        ("P(Net Carry > 0)", result["scenarios"]["positive_carry_probability"]),
    ]
    _write_pairs(ws, 5, carry_pairs, {
        "Net Carry Period Rate", "Net Carry Effective Annual", "Break-even Loan Rate",
        "Break-even Spread", "Safety Margin", "10Y Bond Relative Value", "P(Net Carry > 0)"
    })

    ws = wb.create_sheet(t["curve"])
    _style_sheet(ws, t["curve"])
    headers = ["Tenor", "VKGS", "Maturity Date", "Benchmark Base Rate", "Source Type", "Source", "Node Kind", "Base Credit Spread", "Total Credit Spread", "Liquidity", "Capital", "Fee Eq.", "All-in Quote Rate", "Cash Effective", "After-tax Effective", "Real After-tax"]
    for c, h in enumerate(headers, 1):
        ws.cell(5, c, h)
        ws.cell(5, c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(5, c).font = Font(name="Arial", size=9, bold=False, color=WHITE)
        ws.column_dimensions[get_column_letter(c)].width = 18 if c not in {6} else 34
    for r, row in enumerate(result["curve"]["rows"], 6):
        values = [
            row["label"], row["days"], row["maturity_date"], row.get("base_rate"), row.get("base_source_type"), row.get("base_source"), row.get("node_kind"),
            row.get("base_credit_spread", inp.credit_spread), row["credit_spread"], row["liquidity_premium"], row["capital_charge"], row["fee_equivalent"],
            row.get("all_in_quote_rate"), row.get("cash_effective_annual"), row.get("after_tax_effective_annual"), row.get("real_after_tax_effective")
        ]
        for c, v in enumerate(values, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = Font(name="Arial", size=9, bold=False, color=DARK_TEXT)
            if c in {4, 8, 9, 10, 11, 12, 13, 14, 15, 16} and isinstance(v, (int, float)):
                ws.cell(r, c).number_format = "0.00%"
    chart = LineChart()
    chart.title = "Loan Cost Curve"
    chart.style = 13
    chart.y_axis.title = "Rate"
    chart.x_axis.title = "VKGS"
    data = Reference(ws, min_col=13, max_col=13, min_row=5, max_row=5 + len(result["curve"]["rows"]))
    cats = Reference(ws, min_col=2, min_row=6, max_row=5 + len(result["curve"]["rows"]))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 10
    chart.width = 24
    ws.add_chart(chart, "M5")

    ws = wb.create_sheet(t["ladder"])
    _style_sheet(ws, t["ladder"])
    ladder_headers = ["Tenor", "VKGS", "Maturity Date", "Benchmark Base Rate", "Source Type", "Source", "Node Kind", "Base Credit Spread", "Total Credit Spread", "Liquidity", "Capital", "Fee Eq.", "All-in Quote Rate", "Cash Effective", "After-tax Effective", "Real After-tax", "10Y Bond Spread"]
    for c, h in enumerate(ladder_headers, 1):
        ws.cell(5, c, h)
        ws.cell(5, c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(5, c).font = Font(name="Arial", size=9, bold=False, color=WHITE)
        ws.column_dimensions[get_column_letter(c)].width = 18 if c != 6 else 34
    for r, row in enumerate(result["curve"]["rows"], 6):
        vals = [
            row["label"], row["days"], row["maturity_date"], row.get("base_rate"), row.get("base_source_type"), row.get("base_source"), row.get("node_kind"),
            row.get("base_credit_spread", inp.credit_spread), row["credit_spread"], row["liquidity_premium"], row["capital_charge"], row["fee_equivalent"],
            row.get("all_in_quote_rate"), row.get("cash_effective_annual"), row.get("after_tax_effective_annual"), row.get("real_after_tax_effective"), row.get("ten_year_bond_spread")
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = Font(name="Arial", size=9, bold=False, color=DARK_TEXT)
            if c in {4, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17} and isinstance(v, (int, float)):
                ws.cell(r, c).number_format = "0.00%"

    ws = wb.create_sheet(t["real"])
    _style_sheet(ws, t["real"])
    real_pairs = [
        ("12M Expected CPI (TCMB PKA / Override)", inp.inflation_12m_expectation),
        ("Borrowing Cash Effective Annual", result["borrower"]["cash_effective_annual"]),
        ("Borrowing After-tax Effective Annual", result["borrower"]["after_tax_effective_annual"]),
        ("Borrowing Real After-tax Effective", result["borrower"]["real_after_tax_effective"]),
        ("Investment Gross Effective Annual", result["investment"]["gross_effective_annual"]),
        ("Investment After-tax Effective Annual", result["investment"]["after_tax_effective_annual"]),
        ("Investment Real After-tax Effective", result["investment"]["real_after_tax_effective"]),
        ("10Y TRY Bond Yield", inp.ten_year_bond_yield),
        ("10Y Relative Value", k["ten_year_bond_relative_value"]),
    ]
    _write_pairs(ws, 5, real_pairs, {x[0] for x in real_pairs})

    ws = wb.create_sheet(t["scenarios"])
    _style_sheet(ws, t["scenarios"])
    sheaders = ["Scenario", "Probability", "TLREF", "Credit Spread", "Inflation", "Loan Rate", "Net Carry Amount", "Net Carry Eff.", "Real Carry Eff."]
    for c, h in enumerate(sheaders, 1):
        ws.cell(5, c, h)
        ws.cell(5, c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(5, c).font = Font(name="Arial", size=9, bold=False, color=WHITE)
    for r, row in enumerate(result["scenarios"]["rows"], 6):
        vals = [row["name"], row["probability"], row["tlref"], row["credit_spread"], row["inflation"], row["loan_rate"], row["net_carry_amount"], row["net_carry_effective_annual"], row["real_net_carry_effective"]]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = Font(name="Arial", size=9, bold=False, color=DARK_TEXT)
            if c in {2, 3, 4, 5, 6, 8, 9}:
                ws.cell(r, c).number_format = "0.00%"
            if c == 7:
                ws.cell(r, c).number_format = '#,##0.00;[Red](#,##0.00);-'

    ws = wb.create_sheet(t["sensitivity"])
    _style_sheet(ws, t["sensitivity"])
    sens = result["sensitivity"]
    ws.cell(5, 1, "TLREF / Spread")
    for c, s in enumerate(sens["spread_values"], 2):
        ws.cell(5, c, s)
        ws.cell(5, c).number_format = "0.00%"
    for r, tl in enumerate(sens["tlref_values"], 6):
        ws.cell(r, 1, tl)
        ws.cell(r, 1).number_format = "0.00%"
        for c, v in enumerate(sens["grid"][r - 6], 2):
            ws.cell(r, c, v)
            ws.cell(r, c).number_format = "0.00%"

    ws = wb.create_sheet(t["methodology"])
    _style_sheet(ws, t["methodology"])
    if language == "tr":
        methods = [
            "Tarihsel TLREF motoru: paketlenmiş resmî Borsa İstanbul günlük TLREF CSV serisi kalıcı tarihsel katmandır; webden okunabilen yeni resmî gözlemlerle birleştirilir.",
            "Güncel ankor: son geçerli BIST TLREF gözlemi ON düğümüdür. Tek ON gözlemi 1H-10Y vade yapısını tanımlamaz.",
            "Benchmark kaynak hiyerarşisi: VİOP TLREF futures / swap / kullanıcı piyasa düğümü > TCMB PKA TLREF-eşdeğer anket proxy düğümü > yalnız desteklenen ufuk içinde model interpolasyonu.",
            "TCMB PKA politika-faizi beklentileri, güncel gözlenen (TLREF - TCMB politika faizi) bazı eklenerek TLREF-eşdeğer survey proxy oranına çevrilir. Bu proxy düğümler borsada işlem gören forward/futures fiyatları değildir.",
            "Varsayılan olarak son gerçek benchmark düğümünün ötesine ekstrapolasyon yapılmaz. 10Y TL devlet tahvili, açıkça baz-faiz düğümü yapılmadıkça referans işaretidir.",
            "Nelson-Siegel: y_NS(T) = beta0 + beta1[(1-exp(-T/tau1))/(T/tau1)] + beta2[(1-exp(-T/tau1))/(T/tau1) - exp(-T/tau1)]. En az dört farklı düğüm gerekir.",
            "Svensson: y_SV(T) = y_NS(T) + beta3[(1-exp(-T/tau2))/(T/tau2) - exp(-T/tau2)]. En az altı farklı düğüm gerekir.",
            "Monoton ham ankorlarda NS/Svensson optimizasyonuna şekil cezası eklenir; çizgi parametrik olmayan sonradan düzeltme ile değiştirilmez ve gerçek NS/Svensson fonksiyonu olarak kalır.",
            "Kredi kotasyonu: q_loan(T) = y_benchmark(T) + s_base + s_tenor(T) + l(T) + k(T) + f(T). Spread, likidite, sermaye ve ücretler açık fiyatlama girdileridir.",
            "BSMV, KKDF, ücretler ve vergi kalkanı kotasyon eğrisinden sonra uygulanır; reel maliyet hesabında 12 aylık beklenen TÜFE kullanılır.",
            "Reel maliyet = (1 + vergi sonrası nominal maliyet) / (1 + 12A beklenen TÜFE) - 1.",
            "TLREF gerçekleşen dönem getirisi, günlük resmî seride her oranın geçerli olduğu takvim günleri dikkate alınarak bileşik hesaplanır.",
            "Fit tanıları: ham ankorlar, kaynak tipi, NS/Svensson RMSE, gözlenen düğüm sayısı ve azami gerçek fit ufku saklanır.",
            "Net carry = vergi sonrası yatırım geliri - vergi sonrası kredi maliyeti; başabaş kredi faizi bu iki ekonomik değeri eşitleyen orandır.",
        ]
    else:
        methods = [
            "Historical TLREF engine: the bundled official Borsa Istanbul daily TLREF CSV is the permanent history layer and is merged with newly parsed official web observations.",
            "Current anchor: the latest valid BIST TLREF observation is the ON node. One ON observation does not identify a 1W-10Y term structure.",
            "Benchmark source hierarchy: VIOP TLREF futures / swap / user market node > TCMB PKA TLREF-equivalent survey proxy > model interpolation only inside the supported horizon.",
            "TCMB PKA policy-rate expectations are converted to TLREF-equivalent survey proxy rates by adding the current observed (TLREF - TCMB policy rate) basis. These proxy nodes are never labelled as exchange-traded forward/futures prices.",
            "No extrapolation is performed beyond the last genuine benchmark node by default. The 10Y TRY government bond is a reference marker unless explicitly promoted to a base-rate node.",
            "Nelson-Siegel: y_NS(T) = beta0 + beta1[(1-exp(-T/tau1))/(T/tau1)] + beta2[(1-exp(-T/tau1))/(T/tau1) - exp(-T/tau1)]. At least four distinct nodes are required.",
            "Svensson: y_SV(T) = y_NS(T) + beta3[(1-exp(-T/tau2))/(T/tau2) - exp(-T/tau2)]. At least six distinct nodes are required.",
            "For monotone raw anchors, a shape penalty is included in NS/Svensson optimization; the line remains the actual parametric function and is not replaced by a post-fit non-parametric curve.",
            "Loan quote: q_loan(T) = y_benchmark(T) + s_base + s_tenor(T) + l(T) + k(T) + f(T). Spread, liquidity, capital and fees remain explicit pricing inputs.",
            "BSMV, KKDF, fees and tax shield are applied after the quoted curve; real-cost normalization uses the 12-month expected CPI input.",
            "Real cost = (1 + nominal after-tax cost) / (1 + 12M expected CPI) - 1.",
            "Realized TLREF period return compounds the official daily series using the calendar days for which each daily fixing applies.",
            "Fit diagnostics retain raw anchors, source type, NS/Svensson RMSE, observed-node count and the maximum genuine fit horizon.",
            "Net carry = after-tax investment income - after-tax borrowing cost; the break-even loan rate equalizes those economic values.",
        ]
    for i, line in enumerate(methods, 5):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.cell(i, 1, line)
        ws.cell(i, 1).font = Font(name="Arial", size=10, bold=False, color=DARK_TEXT)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 44

    ws = wb.create_sheet(t["sources"])
    _style_sheet(ws, t["sources"])
    sources = [
        ("GIB Temporary Article 67 Guide 2026", "https://cdn.gib.gov.tr/api/gibportal-file/file/getFile?objectKey=DUYURU%2FUNIVERSAL%2F2026%2F2026_Gecici67.pdf"),
        ("GIB Corporate Tax Return Guide 2026", "https://cdn.gib.gov.tr/api/gibportal-file/file/getFile?objectKey=DUYURU%2FUNIVERSAL%2F2026%2F2026_Kurumlar_Vergisi_Beyan_Rehberi.pdf"),
        ("GIB BSMV consolidated decision", "https://gib.gov.tr/mevzuat/kanun/445/bkk/1849"),
        ("GIB KKDF official ruling page", "https://gib.gov.tr/mevzuat/kanun/445/ozelge/36832"),
        ("Borsa Istanbul TLREF official page / historical series", "https://www.borsaistanbul.com/endeksler/tlref"),
        ("Borsa Istanbul 1 Month TLREF Futures Contract", "https://www.borsaistanbul.com/piyasalar/viop/vadeli-islem-sozlesmeleri/faiz-vadeli-islem-sozlesmeleri"),
        ("TCMB Market Participants Survey", "https://www.tcmb.gov.tr/wps/wcm/connect/TR/TCMB%2BTR/Main%2BMenu/Istatistikler/Egilim%2BAnketleri/Piyasa%2BKatilimcilari%2BAnketi"),
        ("TCMB EVDS - 12M CPI expectation", "https://evds3.tcmb.gov.tr/"),
    ]
    for r, (name, url) in enumerate(sources, 5):
        ws.cell(r, 1, name)
        ws.cell(r, 2, url)
        ws.cell(r, 2).hyperlink = url
        ws.cell(r, 2).style = "Hyperlink"
        ws.cell(r, 1).font = Font(name="Arial", size=9, bold=False, color=DARK_TEXT)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
