from io import BytesIO

from openpyxl import load_workbook

from app.engine import simulate
from app.exporter import build_results_workbook
from app.models import SimulationInput


def _survey_nodes():
    return [
        {"label": "ON", "days": 1, "base_rate": 0.368433, "base_source": "BIST TLREF", "base_source_type": "bist_tlref"},
        {"label": "OCT26", "days": 17, "base_rate": 0.3600, "base_source": "TCMB PKA", "base_source_type": "tcmb_pka"},
        {"label": "DEC26", "days": 66, "base_rate": 0.3507, "base_source": "TCMB PKA", "base_source_type": "tcmb_pka"},
        {"label": "12M", "days": 341, "base_rate": 0.2922, "base_source": "TCMB PKA", "base_source_type": "tcmb_pka"},
        {"label": "DEC27", "days": 452, "base_rate": 0.2716, "base_source": "TCMB PKA", "base_source_type": "tcmb_pka"},
        {"label": "24M", "days": 707, "base_rate": 0.2243, "base_source": "TCMB PKA", "base_source_type": "tcmb_pka"},
    ]


def test_value_only_export_has_no_cell_formulas_and_all_sections():
    inp = SimulationInput(pricing_mode="direct", tlref=0.40, direct_loan_rate=0.40, bsmv_rate=0.05, tenor_points=_survey_nodes())
    result = simulate(inp)
    payload = build_results_workbook(inp, result, "tr")
    wb = load_workbook(BytesIO(payload), data_only=False)
    assert len(wb.sheetnames) == 14
    formulas = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append((ws.title, cell.coordinate))
    assert formulas == []


def test_export_curve_contains_base_rate_and_source_columns():
    inp = SimulationInput(tlref=0.368433, tenor_points=_survey_nodes())
    result = simulate(inp)
    payload = build_results_workbook(inp, result, "en")
    wb = load_workbook(BytesIO(payload), data_only=False)
    ws = wb["Loan Yield Curve"]
    headers = [ws.cell(5, c).value for c in range(1, 17)]
    assert "Benchmark Base Rate" in headers
    assert "Source Type" in headers
    assert "Source" in headers
    assert "All-in Quote Rate" in headers
